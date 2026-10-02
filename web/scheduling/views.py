from io import BytesIO
from urllib.parse import quote

from django.contrib import messages
from django.contrib.auth.decorators import login_required, permission_required
from django.core.exceptions import PermissionDenied
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.views.decorators.cache import never_cache
from django.views.decorators.http import require_POST
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .forms import (
    AssignmentForm,
    AssignmentResponseForm,
    MatchForm,
)
from .models import Assignment, Match, RefereeProfile


@never_cache
@login_required
def home(request):
    context = {
        "has_referee_profile": False,
        "pending_assignments": [],
        "upcoming_assignments": [],
    }

    try:
        referee_profile = request.user.referee_profile
    except RefereeProfile.DoesNotExist:
        pass
    else:
        assignments = Assignment.objects.filter(
            referee=referee_profile,
            match__assignment_status=(
                Match.AssignmentStatus.PUBLISHED
            ),
            match__kickoff_at__date__gte=timezone.localdate(),
        ).select_related(
            "match",
            "match__competition",
            "match__home_team",
            "match__away_team",
            "match__venue",
        ).order_by(
            "match__kickoff_at",
            "position",
        )

        context.update(
            {
                "has_referee_profile": True,
                "pending_assignments": assignments.filter(
                    response_status=(
                        Assignment.ResponseStatus.PENDING
                    )
                ),
                "upcoming_assignments": assignments.exclude(
                    response_status=(
                        Assignment.ResponseStatus.PENDING
                    )
                ),
            }
        )

    return render(
        request,
        "scheduling/home.html",
        context,
    )


@login_required
@permission_required("scheduling.view_match", raise_exception=True)
def match_list(request):
    today = timezone.localdate()

    matches = Match.objects.select_related(
        "competition",
        "home_team",
        "away_team",
        "venue",
    )

    context = {
        "upcoming_matches": matches.filter(
            kickoff_at__date__gte=today
        ).order_by("kickoff_at", "id"),
        "past_matches": matches.filter(
            kickoff_at__date__lt=today
        ).order_by("-kickoff_at", "-id"),
    }

    return render(request, "scheduling/match_list.html", context)


@login_required
@permission_required(
    "scheduling.publish_assignments",
    raise_exception=True,
)
def assignment_export(request):
    matches = Match.objects.filter(
        assignment_status=Match.AssignmentStatus.PUBLISHED
    ).select_related(
        "competition",
        "home_team",
        "away_team",
        "venue",
    ).prefetch_related(
        "assignments__referee",
    ).order_by(
        "kickoff_at",
        "id",
    )

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "裁判安排"
    worksheet.sheet_view.showGridLines = False
    worksheet.freeze_panes = "A4"

    headers = [
        "赛事",
        "赛季",
        "场次编号",
        "轮次",
        "开球时间",
        "主队",
        "客队",
        "比赛场地",
        "赛事级别",
        "比赛状态",
        "主裁判",
        "主裁反馈",
        "第一助理裁判",
        "第一助理反馈",
        "第二助理裁判",
        "第二助理反馈",
        "第四官员",
        "第四官员反馈",
        "安排发布时间",
    ]

    last_column = get_column_letter(len(headers))

    worksheet.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=len(headers),
    )
    title_cell = worksheet["A1"]
    title_cell.value = "足协裁判安排表"
    title_cell.font = Font(
        name="微软雅黑",
        size=16,
        bold=True,
        color="FFFFFF",
    )
    title_cell.fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )
    title_cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
    )
    worksheet.row_dimensions[1].height = 30

    worksheet.merge_cells(
        start_row=2,
        start_column=1,
        end_row=2,
        end_column=len(headers),
    )
    exported_at = timezone.localtime().strftime(
        "%Y年%m月%d日 %H:%M"
    )
    worksheet["A2"] = (
        f"仅包含已经发布的裁判安排　导出时间：{exported_at}"
    )
    worksheet["A2"].font = Font(
        name="微软雅黑",
        size=10,
        color="666666",
    )
    worksheet["A2"].alignment = Alignment(
        horizontal="left",
        vertical="center",
    )
    worksheet.row_dimensions[2].height = 22

    thin_side = Side(
        style="thin",
        color="D9E2F3",
    )
    thin_border = Border(
        left=thin_side,
        right=thin_side,
        top=thin_side,
        bottom=thin_side,
    )

    for column_number, header in enumerate(headers, start=1):
        cell = worksheet.cell(
            row=3,
            column=column_number,
            value=header,
        )
        cell.font = Font(
            name="微软雅黑",
            size=10,
            bold=True,
            color="FFFFFF",
        )
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="4472C4",
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )
        cell.border = thin_border

    worksheet.row_dimensions[3].height = 26

    def local_excel_datetime(value):
        if value is None:
            return None

        if timezone.is_aware(value):
            value = timezone.localtime(value)

        return value.replace(tzinfo=None)

    def assignment_values(assignments, position):
        assignment = assignments.get(position)

        if assignment is None:
            return "", ""

        return (
            assignment.referee.name,
            assignment.get_response_status_display(),
        )

    for match in matches:
        assignments = {
            assignment.position: assignment
            for assignment in match.assignments.all()
        }

        referee_name, referee_status = assignment_values(
            assignments,
            Assignment.Position.REFEREE,
        )
        assistant_1_name, assistant_1_status = assignment_values(
            assignments,
            Assignment.Position.ASSISTANT_1,
        )
        assistant_2_name, assistant_2_status = assignment_values(
            assignments,
            Assignment.Position.ASSISTANT_2,
        )
        fourth_name, fourth_status = assignment_values(
            assignments,
            Assignment.Position.FOURTH_OFFICIAL,
        )

        worksheet.append(
            [
                match.competition.name,
                match.competition.season,
                match.match_number,
                match.round_name,
                local_excel_datetime(match.kickoff_at),
                match.home_team.name,
                match.away_team.name,
                match.venue.name,
                match.match_level,
                match.get_status_display(),
                referee_name,
                referee_status,
                assistant_1_name,
                assistant_1_status,
                assistant_2_name,
                assistant_2_status,
                fourth_name,
                fourth_status,
                local_excel_datetime(match.published_at),
            ]
        )

    for row in worksheet.iter_rows(
        min_row=4,
        max_row=worksheet.max_row,
        min_col=1,
        max_col=len(headers),
    ):
        for cell in row:
            cell.font = Font(
                name="微软雅黑",
                size=10,
            )
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )
            cell.border = thin_border

        row[4].number_format = "yyyy-mm-dd hh:mm"
        row[18].number_format = "yyyy-mm-dd hh:mm"

    column_widths = [
        18,
        14,
        13,
        15,
        18,
        18,
        18,
        18,
        12,
        12,
        14,
        12,
        16,
        12,
        16,
        12,
        14,
        12,
        18,
    ]

    for column_number, width in enumerate(
        column_widths,
        start=1,
    ):
        worksheet.column_dimensions[
            get_column_letter(column_number)
        ].width = width

    worksheet.auto_filter.ref = (
        f"A3:{last_column}{worksheet.max_row}"
    )
    worksheet.print_title_rows = "1:3"
    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True

    output = BytesIO()
    workbook.save(output)

    filename = (
        "裁判安排_"
        f"{timezone.localtime().strftime('%Y%m%d_%H%M')}.xlsx"
    )
    response = HttpResponse(
        output.getvalue(),
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
    )
    response["Content-Disposition"] = (
        "attachment; "
        f"filename*=UTF-8''{quote(filename)}"
    )
    response["Cache-Control"] = "no-store"

    return response


@login_required
@permission_required("scheduling.view_match", raise_exception=True)
def match_detail(request, pk):
    match = get_object_or_404(
        Match.objects.select_related(
            "competition",
            "home_team",
            "away_team",
            "venue",
            "created_by",
        ).prefetch_related(
            "assignments__referee",
        ),
        pk=pk,
    )

    # 草稿阶段只有排班管理员可以看到裁判名单。
    can_view_assignments = (
        match.assignment_status
        == Match.AssignmentStatus.PUBLISHED
        or request.user.has_perm("scheduling.change_assignment")
    )
    assignments_by_position = (
        {
            assignment.position: assignment
            for assignment in match.assignments.all()
        }
        if can_view_assignments
        else {}
    )
    assignment_rows = []

    for position, position_name in Assignment.Position.choices:
        assignment = assignments_by_position.get(position)

        assignment_rows.append(
            {
                "position_name": position_name,
                "assignment": assignment,
                "can_respond": (
                    assignment is not None
                    and match.assignment_status
                    == Match.AssignmentStatus.PUBLISHED
                    and assignment.referee.user_id
                    == request.user.id
                ),
            }
        )

    return render(
        request,
        "scheduling/match_detail.html",
        {
            "match": match,
            "assignment_rows": assignment_rows,
            "can_view_assignments": can_view_assignments,
        },
    )


@login_required
@permission_required("scheduling.add_match", raise_exception=True)
def match_create(request):
    form = MatchForm(
        request.POST if request.method == "POST" else None
    )

    if request.method == "POST" and form.is_valid():
        match = form.save(commit=False)
        match.created_by = request.user
        match.save()

        messages.success(request, "比赛已成功录入。")
        return redirect(
            "scheduling:match_detail",
            pk=match.pk,
        )

    return render(
        request,
        "scheduling/match_form.html",
        {
            "form": form,
            "page_title": "录入比赛",
            "submit_label": "保存比赛",
        },
    )


@login_required
@permission_required("scheduling.change_match", raise_exception=True)
def match_update(request, pk):
    match = get_object_or_404(Match, pk=pk)
    form = MatchForm(
        request.POST if request.method == "POST" else None,
        instance=match,
    )

    if request.method == "POST" and form.is_valid():
        match = form.save()
        messages.success(request, "比赛信息已更新。")

        return redirect(
            "scheduling:match_detail",
            pk=match.pk,
        )

    return render(
        request,
        "scheduling/match_form.html",
        {
            "form": form,
            "match": match,
            "page_title": "修改比赛",
            "submit_label": "保存修改",
        },
    )

@login_required
@permission_required(
    (
        "scheduling.add_assignment",
        "scheduling.change_assignment",
    ),
    raise_exception=True,
)
def assignment_update(request, pk):
    match = get_object_or_404(
        Match.objects.select_related(
            "competition",
            "home_team",
            "away_team",
            "venue",
        ),
        pk=pk,
    )
    form = AssignmentForm(
        request.POST if request.method == "POST" else None,
        match=match,
    )

    if request.method == "POST" and form.is_valid():
        was_published = (
            match.assignment_status
            == Match.AssignmentStatus.PUBLISHED
        )
        changed = form.save(assigned_by=request.user)

        if changed and was_published:
            messages.success(
                request,
                "裁判安排已保存，原安排已退回草稿，"
                "请检查后重新发布。",
            )
        elif changed:
            messages.success(request, "裁判安排已保存。")
        else:
            messages.info(request, "裁判安排没有变化。")

        return redirect(
            "scheduling:match_detail",
            pk=match.pk,
        )

    return render(
        request,
        "scheduling/assignment_form.html",
        {
            "form": form,
            "match": match,
        },
    )

@login_required
@permission_required(
    "scheduling.publish_assignments",
    raise_exception=True,
)
@require_POST
def assignment_publish(request, pk):
    match = get_object_or_404(Match, pk=pk)

    if (
        match.assignment_status
        == Match.AssignmentStatus.PUBLISHED
    ):
        messages.info(request, "本场裁判安排已经发布。")
        return redirect(
            "scheduling:match_detail",
            pk=match.pk,
        )

    required_positions = {
        position
        for position, _label in Assignment.Position.choices
    }
    assigned_positions = set(
        match.assignments.values_list(
            "position",
            flat=True,
        )
    )
    missing_positions = required_positions - assigned_positions

    if missing_positions:
        position_labels = dict(Assignment.Position.choices)
        missing_labels = [
            position_labels[position]
            for position in sorted(missing_positions)
        ]
        messages.error(
            request,
            "还不能发布，以下岗位尚未安排："
            f"{'、'.join(missing_labels)}。",
        )
        return redirect(
            "scheduling:match_detail",
            pk=match.pk,
        )

    match.assignment_status = Match.AssignmentStatus.PUBLISHED
    match.published_at = timezone.now()
    match.save(
        update_fields=[
            "assignment_status",
            "published_at",
            "updated_at",
        ]
    )

    messages.success(request, "裁判安排已经发布。")
    return redirect(
        "scheduling:match_detail",
        pk=match.pk,
    )


@login_required
def assignment_respond(request, pk):
    assignment = get_object_or_404(
        Assignment.objects.select_related(
            "match",
            "match__competition",
            "match__home_team",
            "match__away_team",
            "match__venue",
            "referee",
            "referee__user",
        ),
        pk=pk,
    )

    if assignment.referee.user_id != request.user.id:
        raise PermissionDenied("不能反馈其他裁判的安排。")

    if (
        assignment.match.assignment_status
        != Match.AssignmentStatus.PUBLISHED
    ):
        messages.error(
            request,
            "该场裁判安排尚未发布，暂时不能反馈。",
        )
        return redirect(
            "scheduling:match_detail",
            pk=assignment.match_id,
        )

    form = AssignmentResponseForm(
        request.POST if request.method == "POST" else None,
        instance=assignment,
    )

    if request.method == "POST" and form.is_valid():
        assignment = form.save(commit=False)
        assignment.responded_at = timezone.now()
        assignment.save(
            update_fields=[
                "response_status",
                "response_note",
                "responded_at",
                "updated_at",
            ]
        )

        if (
            assignment.response_status
            == Assignment.ResponseStatus.CONFIRMED
        ):
            messages.success(request, "已经确认参加本场执法。")
        else:
            messages.success(request, "请假申请已经提交。")

        return redirect(
            "scheduling:match_detail",
            pk=assignment.match_id,
        )

    return render(
        request,
        "scheduling/assignment_response_form.html",
        {
            "assignment": assignment,
            "form": form,
        },
    )
