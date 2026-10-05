from datetime import datetime, timedelta
from io import BytesIO, StringIO

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from django.core.management import call_command
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from django.contrib.admin.sites import AdminSite
from openpyxl import load_workbook

from .admin import RefereeProfileAdmin
from .models import (
    Assignment,
    Competition,
    Match,
    RefereeProfile,
    Team,
    Venue,
)


class MatchPagePermissionTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        call_command("setup_roles", stdout=StringIO())

        user_model = get_user_model()

        cls.referee = user_model.objects.create_user(
            username="test_referee",
            password="test-password",
        )
        cls.recorder = user_model.objects.create_user(
            username="test_recorder",
            password="test-password",
        )
        cls.scheduler = user_model.objects.create_user(
            username="test_scheduler",
            password="test-password",
        )
        cls.no_role_user = user_model.objects.create_user(
            username="test_no_role",
            password="test-password",
        )

        cls.referee.groups.add(
            Group.objects.get(name="裁判员")
        )
        cls.recorder.groups.add(
            Group.objects.get(name="场次录入员")
        )
        cls.scheduler.groups.add(
            Group.objects.get(name="排班管理员")
        )

        cls.competition = Competition.objects.create(
            name="测试联赛",
            season="2026测试赛季",
        )
        cls.home_team = Team.objects.create(name="测试主队")
        cls.away_team = Team.objects.create(name="测试客队")
        cls.venue = Venue.objects.create(name="测试场地")

        cls.match = Match.objects.create(
            competition=cls.competition,
            match_number="TEST-001",
            round_name="第一轮",
            kickoff_at=timezone.now() + timedelta(days=1),
            home_team=cls.home_team,
            away_team=cls.away_team,
            venue=cls.venue,
            match_level="测试级别",
            created_by=cls.recorder,
        )

    def form_data(self, **changes):
        data = {
            "competition": self.competition.pk,
            "match_number": "TEST-002",
            "round_name": "第二轮",
            "kickoff_at": timezone.localtime(
                self.match.kickoff_at
            ).strftime("%Y-%m-%dT%H:%M"),
            "home_team": self.home_team.pk,
            "away_team": self.away_team.pk,
            "venue": self.venue.pk,
            "match_level": "测试级别",
            "status": Match.Status.SCHEDULED,
            "notes": "测试备注",
        }
        data.update(changes)
        return data

    def test_anonymous_user_is_redirected_to_login(self):
        response = self.client.get(
            reverse("scheduling:match_list")
        )

        self.assertRedirects(
            response,
            "/accounts/login/?next=/matches/",
            fetch_redirect_response=False,
        )

    def test_user_without_role_receives_403(self):
        self.client.force_login(self.no_role_user)

        response = self.client.get(
            reverse("scheduling:match_list")
        )

        self.assertEqual(response.status_code, 403)

    def test_referee_can_view_but_cannot_change_matches(self):
        self.client.force_login(self.referee)

        list_response = self.client.get(
            reverse("scheduling:match_list")
        )
        detail_response = self.client.get(
            reverse(
                "scheduling:match_detail",
                args=[self.match.pk],
            )
        )
        create_response = self.client.get(
            reverse("scheduling:match_create")
        )
        update_response = self.client.get(
            reverse(
                "scheduling:match_update",
                args=[self.match.pk],
            )
        )

        self.assertEqual(list_response.status_code, 200)
        self.assertEqual(detail_response.status_code, 200)
        self.assertEqual(create_response.status_code, 403)
        self.assertEqual(update_response.status_code, 403)

    def test_editing_roles_can_open_match_forms(self):
        for user in (self.recorder, self.scheduler):
            with self.subTest(username=user.username):
                self.client.force_login(user)

                create_response = self.client.get(
                    reverse("scheduling:match_create")
                )
                update_response = self.client.get(
                    reverse(
                        "scheduling:match_update",
                        args=[self.match.pk],
                    )
                )

                self.assertEqual(
                    create_response.status_code,
                    200,
                )
                self.assertEqual(
                    update_response.status_code,
                    200,
                )

    def test_recorder_can_create_match(self):
        self.client.force_login(self.recorder)
        original_count = Match.objects.count()

        response = self.client.post(
            reverse("scheduling:match_create"),
            self.form_data(),
        )

        self.assertEqual(response.status_code, 302)
        self.assertEqual(
            Match.objects.count(),
            original_count + 1,
        )

        created_match = Match.objects.latest("pk")
        self.assertEqual(
            created_match.created_by,
            self.recorder,
        )
        self.assertEqual(
            created_match.assignment_status,
            Match.AssignmentStatus.DRAFT,
        )

    def test_recorder_can_update_match(self):
        self.client.force_login(self.recorder)

        response = self.client.post(
            reverse(
                "scheduling:match_update",
                args=[self.match.pk],
            ),
            self.form_data(round_name="修改后的轮次"),
        )

        self.assertEqual(response.status_code, 302)

        self.match.refresh_from_db()
        self.assertEqual(
            self.match.round_name,
            "修改后的轮次",
        )
        self.assertEqual(
            self.match.created_by,
            self.recorder,
        )

    def test_same_home_and_away_team_is_rejected(self):
        self.client.force_login(self.recorder)
        original_count = Match.objects.count()

        response = self.client.post(
            reverse("scheduling:match_create"),
            self.form_data(away_team=self.home_team.pk),
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(
            response,
            "主队和客队不能是同一支球队",
        )
        self.assertEqual(
            Match.objects.count(),
            original_count,
        )

class AssignmentPageTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        call_command("setup_roles", stdout=StringIO())

        user_model = get_user_model()
        cls.scheduler = user_model.objects.create_user(
            username="assignment_scheduler",
            password="test-password",
        )
        cls.recorder = user_model.objects.create_user(
            username="assignment_recorder",
            password="test-password",
        )

        cls.scheduler.groups.add(
            Group.objects.get(name="排班管理员")
        )
        cls.recorder.groups.add(
            Group.objects.get(name="场次录入员")
        )

        cls.competition = Competition.objects.create(
            name="排班测试联赛",
            season="2026测试赛季",
        )
        cls.home_team = Team.objects.create(name="排班测试主队")
        cls.away_team = Team.objects.create(name="排班测试客队")
        cls.venue = Venue.objects.create(name="排班测试场地")

        cls.match = Match.objects.create(
            competition=cls.competition,
            kickoff_at=timezone.now() + timedelta(days=1),
            home_team=cls.home_team,
            away_team=cls.away_team,
            venue=cls.venue,
            assignment_status=Match.AssignmentStatus.PUBLISHED,
            published_at=timezone.now(),
            created_by=cls.scheduler,
        )

        cls.referees = []
        for index in range(1, 6):
            user = user_model.objects.create_user(
                username=f"assignment_referee_{index}",
                password="test-password",
            )
            cls.referees.append(
                RefereeProfile.objects.create(
                    user=user,
                    name=f"测试裁判{index}",
                )
            )

        positions = [
            Assignment.Position.REFEREE,
            Assignment.Position.ASSISTANT_1,
            Assignment.Position.ASSISTANT_2,
            Assignment.Position.FOURTH_OFFICIAL,
        ]
        statuses = [
            Assignment.ResponseStatus.CONFIRMED,
            Assignment.ResponseStatus.LEAVE,
            Assignment.ResponseStatus.CONFIRMED,
            Assignment.ResponseStatus.PENDING,
        ]

        for index, position in enumerate(positions):
            Assignment.objects.create(
                match=cls.match,
                referee=cls.referees[index],
                position=position,
                response_status=statuses[index],
                response_note=f"原反馈{index + 1}",
                responded_at=timezone.now(),
                assigned_by=cls.scheduler,
            )

    def form_data(self, **changes):
        data = {
            "referee": self.referees[0].pk,
            "assistant_1": self.referees[1].pk,
            "assistant_2": self.referees[2].pk,
            "fourth_official": self.referees[3].pk,
        }
        data.update(changes)
        return data

    def test_only_scheduler_can_open_assignment_form(self):
        url = reverse(
            "scheduling:assignment_update",
            args=[self.match.pk],
        )

        self.client.force_login(self.recorder)
        recorder_response = self.client.get(url)
        self.assertEqual(recorder_response.status_code, 403)

        self.client.force_login(self.scheduler)
        scheduler_response = self.client.get(url)
        self.assertEqual(scheduler_response.status_code, 200)

    def test_duplicate_referee_is_rejected(self):
        self.client.force_login(self.scheduler)

        response = self.client.post(
            reverse(
                "scheduling:assignment_update",
                args=[self.match.pk],
            ),
            self.form_data(
                assistant_1=self.referees[0].pk,
            ),
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(
            response,
            "该裁判已经被安排为主裁判。",
        )

        self.match.refresh_from_db()
        self.assertEqual(
            self.match.assignment_status,
            Match.AssignmentStatus.PUBLISHED,
        )

    def test_only_changed_position_feedback_is_reset(self):
        unchanged = Assignment.objects.get(
            match=self.match,
            position=Assignment.Position.ASSISTANT_1,
        )
        unchanged_pk = unchanged.pk
        unchanged_responded_at = unchanged.responded_at

        self.client.force_login(self.scheduler)
        response = self.client.post(
            reverse(
                "scheduling:assignment_update",
                args=[self.match.pk],
            ),
            self.form_data(
                referee=self.referees[4].pk,
            ),
        )

        self.assertEqual(response.status_code, 302)

        changed = Assignment.objects.get(
            match=self.match,
            position=Assignment.Position.REFEREE,
        )
        unchanged = Assignment.objects.get(
            match=self.match,
            position=Assignment.Position.ASSISTANT_1,
        )

        self.assertEqual(changed.referee, self.referees[4])
        self.assertEqual(
            changed.response_status,
            Assignment.ResponseStatus.PENDING,
        )
        self.assertEqual(changed.response_note, "")
        self.assertIsNone(changed.responded_at)

        self.assertEqual(unchanged.pk, unchanged_pk)
        self.assertEqual(
            unchanged.response_status,
            Assignment.ResponseStatus.LEAVE,
        )
        self.assertEqual(unchanged.response_note, "原反馈2")
        self.assertEqual(
            unchanged.responded_at,
            unchanged_responded_at,
        )

        self.match.refresh_from_db()
        self.assertEqual(
            self.match.assignment_status,
            Match.AssignmentStatus.DRAFT,
        )
        self.assertIsNone(self.match.published_at)

    def test_referees_can_swap_positions(self):
        self.client.force_login(self.scheduler)

        response = self.client.post(
            reverse(
                "scheduling:assignment_update",
                args=[self.match.pk],
            ),
            self.form_data(
                referee=self.referees[1].pk,
                assistant_1=self.referees[0].pk,
            ),
        )

        self.assertEqual(response.status_code, 302)

        referee_assignment = Assignment.objects.get(
            match=self.match,
            position=Assignment.Position.REFEREE,
        )
        assistant_assignment = Assignment.objects.get(
            match=self.match,
            position=Assignment.Position.ASSISTANT_1,
        )

        self.assertEqual(
            referee_assignment.referee,
            self.referees[1],
        )
        self.assertEqual(
            assistant_assignment.referee,
            self.referees[0],
        )
        self.assertEqual(
            referee_assignment.response_status,
            Assignment.ResponseStatus.PENDING,
        )
        self.assertEqual(
            assistant_assignment.response_status,
            Assignment.ResponseStatus.PENDING,
        )

    def test_complete_assignments_can_be_published(self):
        self.match.assignment_status = (
            Match.AssignmentStatus.DRAFT
        )
        self.match.published_at = None
        self.match.save()

        self.client.force_login(self.scheduler)
        response = self.client.post(
            reverse(
                "scheduling:assignment_publish",
                args=[self.match.pk],
            )
        )

        self.assertEqual(response.status_code, 302)

        self.match.refresh_from_db()
        self.assertEqual(
            self.match.assignment_status,
            Match.AssignmentStatus.PUBLISHED,
        )
        self.assertIsNotNone(self.match.published_at)

    def test_incomplete_assignments_cannot_be_published(self):
        self.match.assignment_status = (
            Match.AssignmentStatus.DRAFT
        )
        self.match.published_at = None
        self.match.save()

        Assignment.objects.filter(
            match=self.match,
            position=Assignment.Position.FOURTH_OFFICIAL,
        ).delete()

        self.client.force_login(self.scheduler)
        response = self.client.post(
            reverse(
                "scheduling:assignment_publish",
                args=[self.match.pk],
            )
        )

        self.assertEqual(response.status_code, 302)

        self.match.refresh_from_db()
        self.assertEqual(
            self.match.assignment_status,
            Match.AssignmentStatus.DRAFT,
        )
        self.assertIsNone(self.match.published_at)

    def test_recorder_cannot_publish_assignments(self):
        self.match.assignment_status = (
            Match.AssignmentStatus.DRAFT
        )
        self.match.published_at = None
        self.match.save()

        self.client.force_login(self.recorder)
        response = self.client.post(
            reverse(
                "scheduling:assignment_publish",
                args=[self.match.pk],
            )
        )

        self.assertEqual(response.status_code, 403)

        self.match.refresh_from_db()
        self.assertEqual(
            self.match.assignment_status,
            Match.AssignmentStatus.DRAFT,
        )

    def test_referee_can_respond_to_own_assignment(self):
        assignment = Assignment.objects.get(
            match=self.match,
            position=Assignment.Position.REFEREE,
        )
        assignment.response_status = (
            Assignment.ResponseStatus.PENDING
        )
        assignment.response_note = ""
        assignment.responded_at = None
        assignment.save()

        self.client.force_login(assignment.referee.user)
        response = self.client.post(
            reverse(
                "scheduling:assignment_respond",
                args=[assignment.pk],
            ),
            {
                "response_status": (
                    Assignment.ResponseStatus.CONFIRMED
                ),
                "response_note": "",
            },
        )

        self.assertEqual(response.status_code, 302)

        assignment.refresh_from_db()
        self.assertEqual(
            assignment.response_status,
            Assignment.ResponseStatus.CONFIRMED,
        )
        self.assertIsNotNone(assignment.responded_at)

    def test_leave_response_requires_note(self):
        assignment = Assignment.objects.get(
            match=self.match,
            position=Assignment.Position.REFEREE,
        )

        self.client.force_login(assignment.referee.user)
        response = self.client.post(
            reverse(
                "scheduling:assignment_respond",
                args=[assignment.pk],
            ),
            {
                "response_status": (
                    Assignment.ResponseStatus.LEAVE
                ),
                "response_note": "",
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(
            response,
            "申请请假时必须填写说明。",
        )

        assignment.refresh_from_db()
        self.assertEqual(
            assignment.response_status,
            Assignment.ResponseStatus.CONFIRMED,
        )

    def test_referee_cannot_respond_for_another_referee(self):
        own_assignment = Assignment.objects.get(
            match=self.match,
            position=Assignment.Position.REFEREE,
        )
        other_assignment = Assignment.objects.get(
            match=self.match,
            position=Assignment.Position.ASSISTANT_1,
        )

        self.client.force_login(own_assignment.referee.user)
        response = self.client.get(
            reverse(
                "scheduling:assignment_respond",
                args=[other_assignment.pk],
            )
        )

        self.assertEqual(response.status_code, 403)

    def test_referee_home_shows_own_pending_assignment(self):
        assignment = Assignment.objects.get(
            match=self.match,
            position=Assignment.Position.REFEREE,
        )
        assignment.response_status = (
            Assignment.ResponseStatus.PENDING
        )
        assignment.response_note = ""
        assignment.responded_at = None
        assignment.save()

        referee_user = assignment.referee.user
        referee_user.groups.add(
            Group.objects.get(name="裁判员")
        )

        self.client.force_login(referee_user)
        response = self.client.get(
            reverse("scheduling:home")
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "我的裁判安排")
        self.assertContains(response, "待确认")
        self.assertContains(response, "排班测试主队")
        self.assertContains(response, "排班测试客队")
        self.assertContains(response, "提交反馈")

    def test_draft_assignment_is_hidden_from_referee_home(self):
        assignment = Assignment.objects.get(
            match=self.match,
            position=Assignment.Position.REFEREE,
        )
        referee_user = assignment.referee.user
        referee_user.groups.add(
            Group.objects.get(name="裁判员")
        )

        self.match.assignment_status = (
            Match.AssignmentStatus.DRAFT
        )
        self.match.published_at = None
        self.match.save()

        self.client.force_login(referee_user)
        response = self.client.get(
            reverse("scheduling:home")
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(
            response,
            "目前没有已发布的近期裁判安排。",
        )
        self.assertNotContains(response, "排班测试主队")

    def test_only_scheduler_can_export_assignments(self):
        export_url = reverse(
            "scheduling:assignment_export"
        )

        self.client.force_login(self.recorder)

        recorder_response = self.client.get(export_url)
        recorder_list_response = self.client.get(
            reverse("scheduling:match_list")
        )

        self.assertEqual(
            recorder_response.status_code,
            403,
        )
        self.assertNotContains(
            recorder_list_response,
            "导出已发布安排",
        )

        self.client.force_login(self.scheduler)

        scheduler_response = self.client.get(export_url)
        scheduler_list_response = self.client.get(
            reverse("scheduling:match_list")
        )

        self.assertEqual(
            scheduler_response.status_code,
            200,
        )
        self.assertEqual(
            scheduler_response["Content-Type"],
            (
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
        )
        self.assertIn(
            ".xlsx",
            scheduler_response["Content-Disposition"],
        )
        self.assertContains(
            scheduler_list_response,
            "导出已发布安排",
        )

    def test_export_contains_published_assignment_data(self):
        self.client.force_login(self.scheduler)

        response = self.client.get(
            reverse("scheduling:assignment_export")
        )

        workbook = load_workbook(
            filename=BytesIO(response.content)
        )

        try:
            worksheet = workbook["裁判安排"]

            self.assertEqual(
                worksheet["A1"].value,
                "足协裁判安排表",
            )

            headers = [
                worksheet.cell(
                    row=3,
                    column=column_number,
                ).value
                for column_number in range(1, 20)
            ]

            self.assertEqual(
                headers,
                [
                    "赛事",
                    "赛季",
                    "场次编号",
                    "轮次",
                    "开球时间",
                    "主队",
                    "客队",
                    "比赛场地",
                    "赛事级别",
                    "比赛状态",
                    "主裁判",
                    "主裁反馈",
                    "第一助理裁判",
                    "第一助理反馈",
                    "第二助理裁判",
                    "第二助理反馈",
                    "第四官员",
                    "第四官员反馈",
                    "安排发布时间",
                ],
            )

            row_values = [
                worksheet.cell(
                    row=4,
                    column=column_number,
                ).value
                for column_number in range(1, 20)
            ]

            self.assertEqual(
                row_values[0],
                "排班测试联赛",
            )
            self.assertEqual(
                row_values[1],
                "2026测试赛季",
            )
            self.assertIsInstance(
                row_values[4],
                datetime,
            )
            self.assertEqual(
                row_values[5],
                "排班测试主队",
            )
            self.assertEqual(
                row_values[6],
                "排班测试客队",
            )
            self.assertEqual(
                row_values[7],
                "排班测试场地",
            )
            self.assertEqual(
                row_values[10],
                "测试裁判1",
            )
            self.assertEqual(
                row_values[11],
                "已经确认",
            )
            self.assertEqual(
                row_values[12],
                "测试裁判2",
            )
            self.assertEqual(
                row_values[13],
                "申请请假",
            )
            self.assertEqual(
                row_values[14],
                "测试裁判3",
            )
            self.assertEqual(
                row_values[15],
                "已经确认",
            )
            self.assertEqual(
                row_values[16],
                "测试裁判4",
            )
            self.assertEqual(
                row_values[17],
                "待确认",
            )
            self.assertIsInstance(
                row_values[18],
                datetime,
            )

            self.assertEqual(
                worksheet["E4"].number_format,
                "yyyy-mm-dd hh:mm",
            )
            self.assertEqual(
                worksheet["S4"].number_format,
                "yyyy-mm-dd hh:mm",
            )
            self.assertEqual(
                worksheet.freeze_panes,
                "A4",
            )
            self.assertEqual(
                worksheet.auto_filter.ref,
                "A3:S4",
            )
        finally:
            workbook.close()

    def test_export_excludes_draft_assignments(self):
        self.match.assignment_status = (
            Match.AssignmentStatus.DRAFT
        )
        self.match.published_at = None
        self.match.save()

        self.client.force_login(self.scheduler)

        response = self.client.get(
            reverse("scheduling:assignment_export")
        )

        workbook = load_workbook(
            filename=BytesIO(response.content)
        )

        try:
            worksheet = workbook["裁判安排"]

            self.assertEqual(
                worksheet.max_row,
                3,
            )
            self.assertNotIn(
                "排班测试主队",
                [
                    cell.value
                    for row in worksheet.iter_rows()
                    for cell in row
                ],
            )
        finally:
            workbook.close()

class RefereeProfileAdminTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        call_command("setup_roles", stdout=StringIO())

        user_model = get_user_model()
        cls.user = user_model.objects.create_user(
            username="new_referee",
            password="test-password",
        )

    def test_saving_profile_adds_referee_group(self):
        profile = RefereeProfile(
            user=self.user,
            name="新建测试裁判",
        )
        model_admin = RefereeProfileAdmin(
            RefereeProfile,
            AdminSite(),
        )

        model_admin.save_model(
            request=None,
            obj=profile,
            form=None,
            change=False,
        )

        self.assertTrue(
            self.user.groups.filter(
                name="裁判员"
            ).exists()
        )
